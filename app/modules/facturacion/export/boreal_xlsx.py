"""Plano de Boreal (O.S. 285 — OSSIMRA - BOREAL) de una factura — SOLO LECTURA.

`GET /api/facturacion/facturas/{id}/export/boreal.xlsx` (Herramientas → Plano Boreal).
Reemplaza al `archivo_plano_2_xls.php` del legacy, con las mismas cuatro columnas:

    NRO. VALIDACIÓN | FECHA REALIZACION PRACTICA | MATRICULA PROVINCIAL | VALOR FACTURADO

- Una fila por prestación del período (`estado <> 'X'`, misma versión de la factura). Los
  ayudantes ya son filas propias del detalle, con su matrícula y su valor.
- Valor facturado = `importe_total`, que ya viene neto de coseguro (el legacy restaba el coseguro
  a las consultas/prácticas). Al pie: TOTAL y COSEGURO (suma de la columna `coseguro`).
- Orden: nombre del médico, tipo de orden (C, P, H, S), fecha de la práctica.

Contra el plano de agosto 2026 (971 filas, $ 18.810.729,90) coinciden 969 filas y el total; las
otras dos son una fecha vacía en la base (el legacy imprimía "30/11/-0001") y un nº de validación
con cero a la izquierda que la base guarda como número.
"""
import datetime
import io
from dataclasses import dataclass
from decimal import Decimal

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import FacturacionCMC

OBRA_SOCIAL = 285
NOMBRE_OBRA_SOCIAL = "OSSIMRA - BOREAL"
TITULO = "COLEGIO MEDICO CORRIENTES"
ENCABEZADOS = ("NRO. VALIDACIÓN", "FECHA REALIZACION PRACTICA", "MATRICULA PROVINCIAL", "VALOR FACTURADO")
ANCHOS = (12.75, 22.25, 17.75, 14.25)

SQL_FILAS = """
    select d.id_detalle_prestaciones as id, d.nro_orden, d.autorizacion, d.fecha_practica,
           d.cod_med, d.importe_total, d.coseguro, d.version,
           m.MATRICULA_PROV as matricula
    from detalle_facturacion d
    left join listado_medico m on m.NRO_SOCIO = d.cod_med
    where d.cod_obr = :os and d.periodo = :periodo and d.estado <> 'X'
    order by m.NOMBRE, field(d.tipo_orden, 'C', 'P', 'H', 'S'), d.fecha_practica, d.id_detalle_prestaciones
"""

_ARIAL = "Arial"
_MEDIO = Side(style="medium")
_BORDE = Border(left=_MEDIO, right=_MEDIO, top=_MEDIO, bottom=_MEDIO)
_CENTRO = Alignment(horizontal="center", vertical="center")
_FORMATO_FECHA = "dd/mm/yyyy"
_FORMATO_MONEDA = "#,##0.00"


@dataclass
class PlanoBoreal:
    contenido: bytes
    nombre_archivo: str
    filas: int
    total: Decimal
    coseguro: Decimal
    sin_fecha: int


def _validacion(fila) -> str:
    nro = str(fila["nro_orden"] or "").strip()
    if nro in ("", "0"):
        nro = (fila["autorizacion"] or "").strip()
    return nro


def _matricula(valor):
    s = str(valor).strip()
    return int(s) if s.isdigit() else s


def _libro(factura: FacturacionCMC, filas: list) -> tuple[bytes, Decimal, Decimal, int]:
    anio, mes = factura.periodo[:4], int(factura.periodo[4:6])
    wb = Workbook()
    ws = wb.active
    ws.title = f"plano boreal {mes:02d}-{anio}"
    ws.sheet_view.showGridLines = False
    for letra, ancho in zip("ABCD", ANCHOS):
        ws.column_dimensions[letra].width = ancho

    ws.merge_cells("A1:D1")
    ws["A1"] = TITULO
    ws["A1"].font = Font(name=_ARIAL, size=11, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", start_color="000000", end_color="000000")
    ws["A1"].alignment = _CENTRO

    ws.merge_cells("A2:D2")
    ws["A2"] = f"OBRA SOCIAL: {NOMBRE_OBRA_SOCIAL} - PERIODO: {mes} - {anio}"
    ws["A2"].font = Font(name=_ARIAL, size=9, bold=True)
    ws["A2"].alignment = _CENTRO
    for letra in "ABCD":
        ws[f"{letra}2"].border = Border(bottom=_MEDIO)

    azul = PatternFill("solid", start_color="0000FF", end_color="0000FF")
    for col, titulo in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=3, column=col, value=titulo)
        c.font = Font(name=_ARIAL, size=8, bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = _CENTRO
        c.border = _BORDE

    fuente = Font(name=_ARIAL, size=8)
    total = coseguro = Decimal(0)
    sin_fecha = 0
    for r in filas:
        fecha = r["fecha_practica"]
        if not isinstance(fecha, datetime.date):
            fecha, sin_fecha = None, sin_fecha + 1
        valor = Decimal(r["importe_total"] or 0)
        total += valor
        coseguro += Decimal(r["coseguro"] or 0)
        ws.append([_validacion(r), fecha, _matricula(r["matricula"]), float(valor)])
        fila = ws.max_row
        for col, formato in enumerate(("@", _FORMATO_FECHA, "General", _FORMATO_MONEDA), start=1):
            c = ws.cell(row=fila, column=col)
            c.font, c.alignment, c.border, c.number_format = fuente, _CENTRO, _BORDE, formato

    negrita = Font(name=_ARIAL, size=8, bold=True)
    for etiqueta, monto in (("TOTAL", total), ("COSEGURO", coseguro)):
        ws.append(["-", "-", etiqueta, float(monto)])
        fila = ws.max_row
        for col in range(1, 5):
            c = ws.cell(row=fila, column=col)
            c.font, c.alignment = negrita, _CENTRO
        ws.cell(row=fila, column=4).number_format = _FORMATO_MONEDA

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), total, coseguro, sin_fecha


async def generar(db: AsyncSession, factura: FacturacionCMC) -> PlanoBoreal:
    """Arma el Excel de `factura` (período cerrado o abierto)."""
    if str(factura.cod_obr).strip() != str(OBRA_SOCIAL):
        raise HTTPException(422, f"El plano de Boreal es sólo de la obra social {OBRA_SOCIAL}.")
    periodo = factura.periodo
    filas = (await db.execute(text(SQL_FILAS), {"os": OBRA_SOCIAL, "periodo": periodo})).mappings().all()
    filas = [f for f in filas if f["version"] == factura.version]
    if not filas:
        raise HTTPException(404, "La factura no tiene prestaciones para exportar.")
    sin_matricula = sorted({f["cod_med"] for f in filas if not f["matricula"]})
    if sin_matricula:
        raise HTTPException(
            422, f"Médicos sin matrícula provincial (nº de socio): {', '.join(map(str, sin_matricula[:20]))}.",
        )
    contenido, total, coseguro, sin_fecha = _libro(factura, filas)
    return PlanoBoreal(
        contenido=contenido,
        nombre_archivo=f"plano boreal {periodo[4:6]}-{periodo[:4]}.xlsx",
        filas=len(filas),
        total=total,
        coseguro=coseguro,
        sin_fecha=sin_fecha,
    )

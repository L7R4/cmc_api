"""Carátula de una factura — PDF + Excel.

Reusa los totales que ya calculó `armado.py` para el detalle (no vuelve a
recorrer las filas): `armado.resumen.por_tipo` sólo trae los tipos con
monto > 0 (regla del RESUMEN GENERAL del detalle), así que acá se lee como
dict y se completa con 0 el que falte — la carátula, a diferencia del detalle,
siempre muestra CONSULTAS/PRACTICAS/HONORARIOS aunque estén en cero; sólo
SANATORIOS se oculta si no hay.

Los dos párrafos fijos y el texto de cierre van tal cual el legacy
(`caratula_pdf.php`/`caratula_excel.php`), decisión del usuario: la variante
"EL PLAZO DE LIQUIDACION..." (la del PDF) en los dos formatos.
"""
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from typing import Optional

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import FacturacionCMC, Institucion, ObrasSociales
from app.modules.facturacion import service
from app.modules.facturacion.export.armado import Armado

PARRAFO_PRESENTACION = (
    "A los efectos de la liquidación y pago correspondiente, cúmplenos remitir "
    "adjunto a las facturas que se detallan seguidamente, por servicios "
    "asistenciales brindados a los afiliados de esa Obra Social. Por los meses "
    "indicados en la misma, juntamente con los comprobantes respectivos:"
)
PARRAFO_PLAZO = (
    "EL PLAZO DE LIQUIDACION DE LA PRESENTE FACTURACION VENCE INDEFECTIBLEMENTE "
    "A LOS 30 (TREINTA) DIAS CORRIDOS DE RECEPCIONADA LA MISMA, VENCIDO DICHO "
    "PLAZO NO SE ACEPTARAN NINGUN TIPO DE DEBITOS Y AUTOMATICAMENTE LA "
    "PRESTATARIA SE COSTITUYE EN MORA, APLICANDOSE DESDE ESE MOMENTO INTERESES "
    "CALCULADOS A LA TASA DEL BANCO NACION (DESCUENTOS DE DOCUMENTOS) VIGENTE "
    "AL DIA DE PAGO DE ESTA LIQUIDACION.-"
)
TEXTO_CIERRE_1 = "En aguardo de esta remesa, saludo a Usted Atentamente."
TEXTO_CIERRE_2 = "SE ADJUNTA FACTURA: 1 original y copias. COMPROBANTES EN TOTAL."

_RAZON_SOCIAL_FALLBACK = "COLEGIO MEDICO DE CORRIENTES"
_CUIT_FALLBACK = "30-57319069-2"
_DOMICILIO_FALLBACK = "CARLOS PELLEGRINI 1785 - (3400) - CORRIENTES"


@dataclass
class MembreteInstitucion:
    razon_social: str
    cuit: str
    domicilio: str
    telefonos: list[str]


@dataclass
class ContextoCaratula:
    membrete: MembreteInstitucion
    nro_expediente: int
    nro_factura: Optional[str]
    periodo_label: str
    obra_social_nombre: str
    cod_obra: str
    estado_label: str
    total_c: Decimal
    total_p: Decimal
    total_h: Decimal
    total_s: Decimal
    subtotal_cp: Decimal
    total_general: Decimal


async def _membrete(db: AsyncSession) -> MembreteInstitucion:
    inst = (await db.execute(select(Institucion).limit(1))).scalars().first()
    if inst is None or not inst.razon_social:
        return MembreteInstitucion(
            razon_social=_RAZON_SOCIAL_FALLBACK, cuit=_CUIT_FALLBACK,
            domicilio=_DOMICILIO_FALLBACK, telefonos=[],
        )
    partes_domicilio = [p for p in (inst.domicilio, inst.localidad, inst.codigo_postal) if p]
    telefonos = [
        f"{t.etiqueta}: {t.numero}" if t.etiqueta else t.numero
        for t in inst.telefonos
    ]
    return MembreteInstitucion(
        razon_social=inst.razon_social,
        cuit=inst.cuit or _CUIT_FALLBACK,
        domicilio=" - ".join(partes_domicilio) or _DOMICILIO_FALLBACK,
        telefonos=telefonos,
    )


async def construir_contexto(
    db: AsyncSession, factura: FacturacionCMC, armado: Armado,
) -> ContextoCaratula:
    totales = dict(armado.resumen.por_tipo)
    total_c = totales.get(service.TIPO_CONSULTA, Decimal("0"))
    total_p = totales.get(service.TIPO_PRACTICA, Decimal("0"))
    total_h = totales.get(service.CATEGORIA_HONORARIOS_INDIVIDUALES, Decimal("0"))
    total_s = totales.get(service.TIPO_SANATORIO, Decimal("0"))

    obra_social_nombre = factura.cod_obr
    try:
        cod_obr_int = int(factura.cod_obr)
    except (TypeError, ValueError):
        cod_obr_int = None
    if cod_obr_int is not None:
        os_row = (await db.execute(
            select(ObrasSociales).where(ObrasSociales.NRO_OBRASOCIAL == cod_obr_int)
        )).scalars().first()
        if os_row is not None:
            obra_social_nombre = os_row.OBRA_SOCIAL

    estado_label = "ABIERTO" if factura.estado == "A" else "CERRADO"

    return ContextoCaratula(
        membrete=await _membrete(db),
        nro_expediente=factura.id_prestaciones,
        nro_factura=factura.nro_factura,
        periodo_label=service.periodo_label(factura.periodo) if factura.periodo else "",
        obra_social_nombre=obra_social_nombre,
        cod_obra=factura.cod_obr,
        estado_label=estado_label,
        total_c=total_c, total_p=total_p, total_h=total_h, total_s=total_s,
        subtotal_cp=total_c + total_p, total_general=armado.resumen.total_general,
    )


# ── PDF ───────────────────────────────────────────────────────────────────

def build_pdf_caratula(ctx: ContextoCaratula) -> bytes:
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 8, ctx.membrete.razon_social, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 5, f"CUIT: {ctx.membrete.cuit}", align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    pdf.ln(2)
    pdf.cell(0, 5, ctx.membrete.domicilio, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    for tel in ctx.membrete.telefonos:
        pdf.cell(0, 5, tel, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(8)

    x_box, box_w = (210 - 170) / 2, 170
    pdf.set_x(x_box)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(box_w, 6, f"Nro Expediente : {ctx.nro_expediente}", align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    y_box, line_h = pdf.get_y(), 7
    pdf.rect(x_box, y_box, box_w, line_h * 3)
    pdf.line(x_box, y_box + line_h, x_box + box_w, y_box + line_h)
    pdf.line(x_box, y_box + line_h * 2, x_box + box_w, y_box + line_h * 2)

    def _campo(x, y, w_label, label, w_val, valor):
        pdf.set_xy(x, y)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(w_label, line_h, label, new_x=XPos.RIGHT, new_y=YPos.TOP)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(w_val, line_h, valor, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    _campo(x_box + 2, y_box, 25, "FACTURA:", 50, ctx.nro_factura or "-")
    _campo(x_box + 90, y_box, 25, "PERIODO:", 50, ctx.periodo_label)
    _campo(x_box + 2, y_box + line_h, 30, "OBRA SOCIAL:", 130, f"({ctx.cod_obra}) - {ctx.obra_social_nombre}")
    _campo(x_box + 2, y_box + line_h * 2, 20, "ESTADO:", 50, ctx.estado_label)

    pdf.set_xy(x_box, y_box + line_h * 3 + 8)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(box_w, 5, PARRAFO_PRESENTACION, align="J")
    pdf.ln(6)

    col1_w, col2_w, h_row = 110, 50, 8
    pdf.set_x(x_box)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(col1_w, h_row, "CONCEPTO", border=1, align="C")
    pdf.cell(col2_w, h_row, "IMPORTE", border=1, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    def _fila_tabla(label, monto, negrita=False):
        pdf.set_x(x_box)
        pdf.set_font("Helvetica", "B" if negrita else "", 10)
        pdf.cell(col1_w, h_row, f"  {label}", border=1, align="L")
        pdf.cell(col2_w, h_row, f"{monto:,.2f}  ", border=1, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    _fila_tabla("CONSULTAS HONORARIOS MEDICOS", ctx.total_c)
    _fila_tabla("PRACTICAS HONORARIOS MEDICOS", ctx.total_p)
    _fila_tabla("SUB TOTAL CONSULTA - PRACTICA", ctx.subtotal_cp, negrita=True)
    _fila_tabla("HONORARIOS INDIVIDUAL", ctx.total_h)
    if ctx.total_s > 0:
        _fila_tabla("SANATORIOS", ctx.total_s)

    pdf.ln(2)
    pdf.set_x(x_box)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(col1_w, 10, "  TOTAL PRESTACIONES", border=1, align="L")
    pdf.cell(col2_w, 10, f"{ctx.total_general:,.2f}  ", border=1, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.ln(15)
    pdf.set_x(x_box)
    pdf.set_font("Helvetica", "", 9)
    pdf.multi_cell(box_w, 5, PARRAFO_PLAZO, align="J")

    pdf.ln(10)
    pdf.set_x(x_box)
    pdf.multi_cell(box_w, 5, f"{TEXTO_CIERRE_1}\n{TEXTO_CIERRE_2}", align="L")

    return bytes(pdf.output())


# ── Excel ─────────────────────────────────────────────────────────────────

def build_excel_caratula(ctx: ContextoCaratula) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Carátula"

    negrita = Font(bold=True)
    negrita_grande = Font(bold=True, size=14)
    centrado = Alignment(horizontal="center", wrap_text=True)

    ws.merge_cells("A1:C1")
    ws["A1"] = ctx.membrete.razon_social
    ws["A1"].font = negrita_grande
    ws["A1"].alignment = centrado

    ws.merge_cells("A2:C2")
    ws["A2"] = f"FACTURA: {ctx.nro_factura or '-'}"
    ws["A2"].font = negrita
    ws["A2"].alignment = centrado

    ws.merge_cells("A3:C3")
    ws["A3"] = f"OBRA SOCIAL: ({ctx.cod_obra}) - {ctx.obra_social_nombre}"
    ws["A3"].font = negrita
    ws["A3"].alignment = centrado

    ws.merge_cells("A5:C6")
    ws["A5"] = PARRAFO_PRESENTACION
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")

    fila = 8
    ws.cell(fila, 1, "CONSULTAS HONORARIOS MEDICOS:").font = negrita
    ws.cell(fila, 3, float(ctx.total_c)).number_format = "#,##0.00"
    fila += 1
    ws.cell(fila, 1, "PRACTICAS HONORARIOS MEDICOS:").font = negrita
    ws.cell(fila, 3, float(ctx.total_p)).number_format = "#,##0.00"
    fila += 1
    ws.cell(fila, 1, "SUB TOTAL CONSULTA - PRACTICA:").font = negrita
    ws.cell(fila, 3, float(ctx.subtotal_cp)).number_format = "#,##0.00"
    fila += 1
    ws.cell(fila, 1, "HONORARIOS INDIVIDUAL:").font = negrita
    ws.cell(fila, 3, float(ctx.total_h)).number_format = "#,##0.00"
    fila += 1
    if ctx.total_s > 0:
        ws.cell(fila, 1, "SANATORIOS:").font = negrita
        ws.cell(fila, 3, float(ctx.total_s)).number_format = "#,##0.00"
        fila += 1
    ws.cell(fila, 1, "TOTAL PRESTACIONES:").font = negrita
    ws.cell(fila, 3, float(ctx.total_general)).number_format = "#,##0.00"
    fila += 2

    ws.merge_cells(f"A{fila}:C{fila + 2}")
    ws.cell(fila, 1, PARRAFO_PLAZO).alignment = Alignment(wrap_text=True, vertical="top")
    fila += 4

    ws.cell(fila, 1, TEXTO_CIERRE_1)
    fila += 1
    ws.cell(fila, 1, TEXTO_CIERRE_2)

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()

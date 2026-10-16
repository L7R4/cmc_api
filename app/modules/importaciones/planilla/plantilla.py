"""La plantilla .xlsx que se descarga en la pantalla de importación.

Se arma desde `columnas.COLUMNAS`, así que nunca se desincroniza con lo que el
servicio espera. Dos hojas: «Prestaciones» (vacía, con listas desplegables) e
«Instrucciones» (qué va en cada columna).
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.modules.importaciones.planilla.columnas import (
    CALCULOS, CLAVES, COLUMNAS, HOJA, ROLES, TIPOS, VIAS,
)

_AZUL = PatternFill("solid", fgColor="1F4E78")
_ROJO = PatternFill("solid", fgColor="C00000")
_BLANCO = Font(bold=True, color="FFFFFF")
_FINO = Side(style="thin", color="BFBFBF")
_FILAS_CON_VALIDACION = 2000
_FILAS_CON_FORMATO_TEXTO = 500

REGLAS = (
    "La obra social y el período se eligen en la pantalla al subir el archivo; no van en la planilla.",
    "Una fila = una prestación. No se agregan filas de totales, subtotales, títulos ni separadores.",
    "Las columnas marcadas con * (en rojo) son obligatorias.",
    "Los importes son POR UNIDAD: el total es (honorarios + gastos + ayudante − coseguro) × cantidad × sesión.",
    "Equipo quirúrgico: una fila por integrante con la misma clave en EQUIPO. El cirujano lleva ROL MEDICO y "
    "honorarios; el ayudante, ROL AYUDANTE y su importe en AYUDANTE (con honorarios y gastos en 0); el pediatra, "
    "ROL PEDIATRA con su propio código y honorarios. Un equipo entra entero o no entra.",
    "SANATORIO: NRO_SOCIO es el médico que ejecuta y CLINICA la clínica que factura. "
    "HONORARIOS INDIVIDUALES: la CLINICA es sólo el lugar.",
    "Antes de grabar se revisa todo fila por fila: lo que no entra se informa con el motivo y no se graba.",
)


def _titulo(clave: str, obligatoria: bool) -> str:
    return clave + (" *" if obligatoria else "")


def generar_plantilla() -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = HOJA
    for j, c in enumerate(COLUMNAS, 1):
        cel = ws.cell(row=1, column=j, value=_titulo(c.clave, c.obligatoria))
        cel.fill = _ROJO if c.obligatoria else _AZUL
        cel.font = _BLANCO
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.comment = Comment(c.descripcion, "CMC")
        ws.column_dimensions[get_column_letter(j)].width = c.ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}1"

    for clave, opciones in (("TIPO", TIPOS), ("VIA", VIAS), ("CALCULO", CALCULOS), ("ROL", ROLES)):
        letra = get_column_letter(CLAVES.index(clave) + 1)
        dv = DataValidation(
            type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True,
            showErrorMessage=True, errorTitle=clave, error="Valor fuera de la lista.",
        )
        dv.add(f"{letra}2:{letra}{_FILAS_CON_VALIDACION}")
        ws.add_data_validation(dv)
    for clave in ("CANTIDAD", "SESION", "PORCENTAJE"):
        letra = get_column_letter(CLAVES.index(clave) + 1)
        dv = DataValidation(
            type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
            showErrorMessage=True, error="Entero mayor o igual a 1.",
        )
        dv.add(f"{letra}2:{letra}{_FILAS_CON_VALIDACION}")
        ws.add_data_validation(dv)
    # Formato de celda por columna: el código, la autorización y el afiliado son texto
    # (si no, Excel les come los ceros a la izquierda).
    for j, c in enumerate(COLUMNAS, 1):
        for i in range(2, _FILAS_CON_FORMATO_TEXTO + 1):
            ws.cell(row=i, column=j).number_format = c.formato

    ins = wb.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 16
    ins.column_dimensions["B"].width = 13
    ins.column_dimensions["C"].width = 95
    ins.column_dimensions["D"].width = 30
    ins["A1"] = "Plantilla de importación de prestaciones"
    ins["A1"].font = Font(bold=True, size=14)
    for k, regla in enumerate(REGLAS, 3):
        ins.cell(row=k, column=1, value=f"{k - 2}. {regla}")
        ins.merge_cells(start_row=k, start_column=1, end_row=k, end_column=4)
        ins.cell(row=k, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ins.row_dimensions[k].height = 32
    base = len(REGLAS) + 4
    for j, t in enumerate(("Columna", "Obligatoria", "Descripción", "Ejemplo"), 1):
        cel = ins.cell(row=base, column=j, value=t)
        cel.fill, cel.font = _AZUL, _BLANCO
    for k, c in enumerate(COLUMNAS, base + 1):
        ins.cell(row=k, column=1, value=c.clave).font = Font(bold=True)
        ins.cell(row=k, column=2, value="Sí" if c.obligatoria else "No")
        ins.cell(row=k, column=3, value=c.descripcion).alignment = Alignment(wrap_text=True, vertical="top")
        ins.cell(row=k, column=4, value=c.ejemplo)
        for j in range(1, 5):
            ins.cell(row=k, column=j).border = Border(top=_FINO, bottom=_FINO, left=_FINO, right=_FINO)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

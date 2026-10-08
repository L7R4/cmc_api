"""Export con la misma estructura y orden que la vista del listado.

Puro (sin DB): arma `FilaExport` a mano y revisa la estructura de `armar`, y que
Excel y PDF se generen con ella.
"""
import dataclasses
import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.modules.facturacion.export.armado import armar
from app.modules.facturacion.export.datos import FilaExport
from app.modules.facturacion.export.encabezado import EncabezadoExport
from app.modules.facturacion.export.excel import build_excel_detalle
from app.modules.facturacion.export.pdf import build_pdf_detalle
from app.modules.facturacion.export.schemas import ExportOpciones


def _fila(id_, cod, nombre, tipo, fecha, afiliado, *, equipo=None, monto="100", created=None,
          clinica=None, clinica_nombre=None):
    return FilaExport(
        id=id_, cod_medico=cod, prestador_nombre=nombre, cod_obr="411", obra_social_nombre=None,
        matricula=None, autorizacion=None, fecha_practica=datetime.date(2026, 9, fecha),
        codigo="420101", nro_afiliado=None, afiliado=afiliado, cantidad=1, sesion=1, porcentaje=100,
        honorarios=Decimal(monto), gastos=Decimal("0"), ayudante=Decimal("0"), coseguro=Decimal("0"),
        subtotal=Decimal(monto), tipo=tipo, tipo_prestador=None, diagnostico=None, via=None,
        especialidad_nombre=None, estado_validacion=None, revisado=False, grupo_equipo_id=equipo,
        created=created, cod_clinica=clinica, clinica_nombre=clinica_nombre,
    )


def _filas():
    filas = [
        _fila(1, "20", "ZAPATA, JUAN", "Consulta", 3, "B"),
        _fila(2, "10", "ACOSTA, MARIA", "Practica", 4, "SOSA"),
        _fila(3, "10", "ACOSTA, MARIA", "Consulta", 2, "RUIZ"),
        _fila(4, "10", "ACOSTA, MARIA", "Honorarios individuales", 12, "PEREYRA"),
        _fila(5, "10", "ACOSTA, MARIA", "Consulta", 9, "BENITEZ"),
        _fila(6, "10", "ACOSTA, MARIA", "Practica", 18, "ALVAREZ"),
        _fila(7, "10", "ACOSTA, MARIA", "Honorarios individuales", 22, "DIAZ"),
        _fila(8, "90", "SOSA, ANA", "Sanatorio", 15, "ARCE", clinica=501, clinica_nombre="DEL SUR"),
        _fila(9, "80", "LEIVA, PABLO", "Sanatorio", 1, "LOPEZ", clinica=502, clinica_nombre="MODELO"),
        _fila(10, "80", "LEIVA, PABLO", "Sanatorio", 11, "CASTRO", clinica=502, clinica_nombre="MODELO"),
        _fila(12, "80", "LEIVA, PABLO", "Sanatorio", 5, "BRAVO", clinica=501, clinica_nombre="DEL SUR"),
        # Ayudante (otro socio) de la práctica 6: va ÚNICAMENTE bajo su cabeza.
        _fila(11, "30", "BRAVO, LUIS", "Practica", 18, "ALVAREZ", equipo=6, monto="20"),
    ]
    filas[5].grupo_equipo_id = 6
    return filas


def _armado(agrupacion="por_socio", **kw):
    return armar(_filas(), ExportOpciones(agrupacion=agrupacion, **kw))


def _grupos(a):
    return {g.cod_medico: g for g in a.secciones[0].grupos}


def test_por_socio_orden_fijo_y_equipo_solo_bajo_su_cabeza():
    a = _armado(orden="importe_desc")
    # Los médicos y, al final, el bloque único de sanatorios.
    assert [s.titulo for s in a.secciones] == [None, "SANATORIOS"]
    grupos = a.secciones[0].grupos

    # Socios A-Z (ignora `orden`); los que solo tienen sanatorios (80, 90) no llevan grupo.
    # El ayudante (30) no tiene grupo propio: va bajo su cabeza.
    assert [g.cod_medico for g in grupos] == ["10", "20"]
    acosta = grupos[0]
    assert acosta.titulo == "10 - ACOSTA, MARIA"
    # Consultas y prácticas: más nueva primero; honorarios: paciente A-Z.
    assert [l.fila.id for l in acosta.lineas] == [5, 3, 6, 2, 7, 4]
    assert [l.subtitulo for l in acosta.lineas] == [
        "CONSULTAS", None, "PRACTICAS", None, "HONORARIOS INDIVIDUALES", None,
    ]
    assert [h.id for h in acosta.lineas[2].hijos] == [11]
    assert acosta.total_general == Decimal("620.00")  # incluye al ayudante, una sola vez

    assert a.una_hoja
    assert a.total_prestaciones == 12
    assert a.resumen.total_general == Decimal("1120.00")
    assert sum(g.total_general for s in a.secciones for g in s.grupos) == a.resumen.total_general


def test_sanatorios_van_en_un_bloque_clinica_paciente_socio():
    [sanatorios] = _armado().secciones[1].grupos
    # DEL SUR: ARCE (socio 90), BRAVO (socio 80) · MODELO: CASTRO, LOPEZ — de cualquier socio.
    assert [l.fila.id for l in sanatorios.lineas] == [8, 12, 10, 9]
    assert [l.subtitulo_clinica for l in sanatorios.lineas] == ["CLINICA DEL SUR", None, "CLINICA MODELO", None]
    # Paciente y socio solo ordenan: sin subtítulo ni total por paciente.
    assert all(l.subtitulo_paciente is None and l.total_paciente is None for l in sanatorios.lineas)
    assert not sanatorios.mostrar_resumen


def test_sanatorios_mismo_paciente_en_la_clinica_va_por_socio_a_z():
    filas = [
        _fila(1, "30", "ZAPATA, JUAN", "Sanatorio", 2, "RUIZ", clinica=501, clinica_nombre="DEL SUR"),
        _fila(2, "10", "ACOSTA, MARIA", "Sanatorio", 9, "RUIZ", clinica=501, clinica_nombre="DEL SUR"),
        _fila(3, "20", "MORENO, ANA", "Sanatorio", 5, "RUIZ", clinica=501, clinica_nombre="DEL SUR"),
        _fila(4, "10", "ACOSTA, MARIA", "Sanatorio", 1, "AGUIRRE", clinica=501, clinica_nombre="DEL SUR"),
        _fila(5, "10", "ACOSTA, MARIA", "Sanatorio", 1, "AGUIRRE", clinica=502, clinica_nombre="ARCOIRIS"),
    ]
    for agrupacion in ("por_socio", "por_tipo"):
        a = armar(filas, ExportOpciones(agrupacion=agrupacion, orden_sanatorio="paciente"))
        [sanatorios] = a.secciones[-1].grupos
        # ARCOIRIS primero (clínica A-Z); luego DEL SUR: AGUIRRE y RUIZ, este por socio ACOSTA, MORENO, ZAPATA.
        assert [l.fila.id for l in sanatorios.lineas] == [5, 4, 2, 3, 1], agrupacion


def test_agrupar_equipo_apagado_se_ignora_el_equipo_siempre_va_con_su_cabeza():
    a = _armado(agrupar_equipo=False)
    grupos = _grupos(a)
    assert "30" not in grupos
    assert [h.id for h in grupos["10"].lineas[2].hijos] == [11]
    assert grupos["10"].total_general == Decimal("620.00")


def test_cabeza_fuera_de_filtro_se_lleva_a_todo_su_equipo():
    filas = _filas()
    filas[5] = dataclasses.replace(filas[5], fuera_de_filtro=True)  # la cabeza (id 6) quedó afuera de los filtros
    a = armar(filas, ExportOpciones(agrupacion="por_socio"))
    grupos = _grupos(a)
    assert "30" not in grupos  # el ayudante no queda suelto: sigue a su cabeza
    assert grupos["10"].total_general == Decimal("500.00")
    assert a.total_prestaciones == 10


def test_equipo_filtrado_igual_va_completo_bajo_la_cabeza():
    filas = _filas()
    filas[11] = dataclasses.replace(filas[11], fuera_de_filtro=True)  # el ayudante (id 11) quedó afuera del filtro
    a = armar(filas, ExportOpciones(agrupacion="por_socio"))
    grupos = _grupos(a)
    assert "30" not in grupos
    assert [h.id for h in grupos["10"].lineas[2].hijos] == [11]
    assert a.total_prestaciones == 12


def test_por_tipo_subtotal_por_medico_y_orden_con_direccion():
    a = _armado("por_tipo", orden="fecha_practica", direccion="desc")
    assert [s.titulo for s in a.secciones] == ["CONSULTA", "PRACTICA", "HONORARIO", "SANATORIO"]
    consultas = a.secciones[0]
    assert [g.cod_medico for g in consultas.grupos] == ["10", "20"]  # socios A-Z
    assert all(g.subtotal_medico and not g.mostrar_resumen for g in consultas.grupos)
    assert [l.fila.id for l in consultas.grupos[0].lineas] == [5, 3]  # fecha, de la más nueva a la más vieja
    assert consultas.grupos[0].total_general == Decimal("200.00")

    # El ayudante cuenta en el subtotal de su cabeza (sección de Prácticas) y no tiene sección propia.
    practicas = a.secciones[1]
    assert [g.cod_medico for g in practicas.grupos] == ["10"]
    assert practicas.grupos[0].total_general == Decimal("220.00")

    asc = _armado("por_tipo", orden="fecha_practica", direccion="asc")
    assert [l.fila.id for l in asc.secciones[0].grupos[0].lineas] == [3, 5]

    # Sanatorio en por_tipo: clínicas seguidas con su subtítulo, dentro de cada socio.
    sanatorios = {g.cod_medico: g for g in a.secciones[3].grupos}
    assert [l.subtitulo_clinica for l in sanatorios["80"].lineas] == ["CLINICA DEL SUR", "CLINICA MODELO", None]


def test_plana_ordena_por_fecha_de_carga_en_ambos_sentidos():
    base = datetime.datetime(2026, 10, 3, 9, 0)
    filas = _filas()[:4]
    for fila, minutos in zip(filas, (30, 10, 20, 20)):  # ids 3 y 4 empatan: desempata el id
        fila.created = base + datetime.timedelta(minutes=minutos)

    def ids(direccion):
        a = armar(filas, ExportOpciones(agrupacion="plana", orden="fecha_carga", direccion=direccion))
        return [l.fila.id for l in a.secciones[0].grupos[0].lineas]

    assert ids("asc") == [2, 3, 4, 1]
    assert ids("desc") == [1, 4, 3, 2]


def test_excel_en_una_sola_hoja_y_pdf_se_generan():
    a = _armado()
    opciones = ExportOpciones(agrupacion="por_socio")
    enc = EncabezadoExport(lineas=["CMC", "OS 411"])
    wb = load_workbook(BytesIO(build_excel_detalle(a, opciones, enc)))
    assert wb.sheetnames == ["Detalle"]
    textos = [r[0] for r in wb["Detalle"].iter_rows(values_only=True) if r and r[0]]
    for esperado in ("10 - ACOSTA, MARIA", "CONSULTAS", "SANATORIOS",
                     "CLINICA MODELO", "CLINICA DEL SUR"):
        assert esperado in textos
    assert not any("CLINICAS / SANATORIOS" == t for t in textos)
    assert any(str(t).startswith("SUBTOTAL SANATORIOS") for t in textos)  # cierra el bloque único
    assert build_pdf_detalle(a, opciones, enc).startswith(b"%PDF")


def test_por_tipo_excel_y_pdf_llevan_el_subtotal_de_cada_medico():
    a = _armado("por_tipo")
    opciones = ExportOpciones(agrupacion="por_tipo")
    enc = EncabezadoExport(lineas=["CMC", "OS 411"])
    wb = load_workbook(BytesIO(build_excel_detalle(a, opciones, enc)))
    assert wb.sheetnames[:2] == ["CONSULTA", "PRACTICA"]
    textos = [r[0] for r in wb["CONSULTA"].iter_rows(values_only=True) if r and r[0]]
    assert "SUBTOTAL SOCIO 10 ACOSTA, MARIA (2): 200.00" in textos
    assert "SUBTOTAL SOCIO 20 ZAPATA, JUAN (1): 100.00" in textos
    assert build_pdf_detalle(a, opciones, enc).startswith(b"%PDF")


def test_por_tipo_orden_por_paciente_en_honorarios_y_sanatorios():
    base = dict(orden="fecha_practica", direccion="asc")

    # Por defecto ("medico"): el orden de siempre (acá, por fecha de práctica).
    a = _armado("por_tipo", **base)
    honorarios = {g.cod_medico: g for g in a.secciones[2].grupos}
    assert [l.fila.id for l in honorarios["10"].lineas] == [4, 7]  # PEREYRA (12), DIAZ (22)
    sanatorios = {g.cod_medico: g for g in a.secciones[3].grupos}
    assert [l.fila.id for l in sanatorios["80"].lineas] == [12, 9, 10]  # DEL SUR · MODELO: por fecha

    # "paciente": la sección entera por paciente A-Z, sin partir por socio; en Sanatorios,
    # dentro de cada clínica, que siguen agrupadas A-Z (con su subtítulo una sola vez).
    b = _armado("por_tipo", orden_honorarios="paciente", orden_sanatorio="paciente", **base)
    [honorarios] = b.secciones[2].grupos
    assert [l.fila.id for l in honorarios.lineas] == [7, 4]  # DIAZ, PEREYRA
    [sanatorios] = b.secciones[3].grupos
    # DEL SUR: ARCE (socio 90), BRAVO (socio 80) · MODELO: CASTRO, LOPEZ
    assert [l.fila.id for l in sanatorios.lineas] == [8, 12, 10, 9]
    assert [l.subtitulo_clinica for l in sanatorios.lineas] == ["CLINICA DEL SUR", None, "CLINICA MODELO", None]
    assert not sanatorios.subtotal_medico

    # Cada selector es independiente y no toca las demás secciones.
    c = _armado("por_tipo", orden_honorarios="paciente", **base)
    assert [l.fila.id for l in {g.cod_medico: g for g in c.secciones[3].grupos}["80"].lineas] == [12, 9, 10]
    assert [l.fila.id for l in c.secciones[0].grupos[0].lineas] == [3, 5]  # Consultas: sin cambios


# ── Honorarios individuales / Sanatorios: pacientes, equipo que no suma, ayudantes sueltos ──


def _filas_hi():
    """Socio 10 con dos cirugías de Honorarios individuales (pacientes DIAZ y PEREYRA, la de
    DIAZ con ayudante y pediatra de otros socios) y una consulta con ayudante."""
    filas = [
        _fila(1, "10", "ACOSTA, MARIA", "Honorarios individuales", 12, "PEREYRA", monto="1000"),
        _fila(2, "10", "ACOSTA, MARIA", "Honorarios individuales", 22, "DIAZ", monto="500"),
        _fila(3, "30", "BRAVO, LUIS", "Honorarios individuales", 22, "DIAZ", monto="100", equipo=2),
        _fila(4, "40", "CORTES, ANA", "Honorarios individuales", 22, "DIAZ", monto="50", equipo=2),
        _fila(5, "10", "ACOSTA, MARIA", "Consulta", 3, "RUIZ", monto="10"),
        _fila(6, "30", "BRAVO, LUIS", "Consulta", 3, "RUIZ", monto="5", equipo=5),
    ]
    filas[1].grupo_equipo_id = 2
    filas[4].grupo_equipo_id = 5
    return filas


def test_por_socio_pacientes_con_subtitulo_y_el_equipo_de_honorarios_no_suma():
    a = armar(_filas_hi(), ExportOpciones(agrupacion="por_socio"))
    g = {x.cod_medico: x for x in a.secciones[0].grupos}["10"]
    hi = [l for l in g.lineas if l.fila.tipo == "Honorarios individuales"]
    assert [l.fila.id for l in hi] == [2, 1]  # DIAZ, PEREYRA
    assert [l.subtitulo_paciente for l in hi] == ["PACIENTE DIAZ", "PACIENTE PEREYRA"]
    assert [h.id for h in hi[0].hijos] == [3, 4]  # el equipo se muestra bajo su cabeza...
    # ...pero no suma: 500 + 1000 de honorarios, y 10 + 5 de la consulta (ahí el equipo sí suma).
    assert g.stats_por_tipo["Honorarios individuales"].monto == Decimal("1500")
    assert g.stats_por_tipo["Consulta"].monto == Decimal("15")
    assert g.total_general == Decimal("1515.00")
    # El resumen general cierra con todo lo facturado.
    assert a.resumen.total_general == Decimal("1665.00")
    # Cada paciente cierra con su total (con el equipo): DIAZ 500 + 100 + 50, PEREYRA 1000.
    assert [l.total_paciente for l in hi] == [("DIAZ", Decimal("650.00")), ("PEREYRA", Decimal("1000.00"))]
    assert all(l.total_paciente is None for l in g.lineas if l.fila.tipo == "Consulta")


def test_por_tipo_honorarios_por_medico_no_suma_el_equipo():
    a = armar(_filas_hi(), ExportOpciones(agrupacion="por_tipo"))
    hi = next(s for s in a.secciones if s.titulo == "HONORARIO").grupos[0]
    assert hi.subtotal_medico and hi.total_general == Decimal("1500.00")
    assert all(l.subtitulo_paciente is None and l.total_paciente is None for l in hi.lineas)


def test_por_tipo_por_paciente_reemplaza_el_subtotal_por_el_total_de_cada_paciente():
    a = armar(_filas_hi(), ExportOpciones(agrupacion="por_tipo", orden_honorarios="paciente"))
    hi = next(s for s in a.secciones if s.titulo == "HONORARIO").grupos[0]
    assert not hi.subtotal_medico
    assert [l.subtitulo_paciente for l in hi.lineas] == ["PACIENTE DIAZ", "PACIENTE PEREYRA"]
    # Total de cada paciente: con el equipo.
    assert [l.total_paciente for l in hi.lineas] == [("DIAZ", Decimal("650.00")), ("PEREYRA", Decimal("1000.00"))]
    assert hi.total_general == Decimal("1500.00")

    enc = EncabezadoExport(lineas=["CMC", "OS 411"])
    opciones = ExportOpciones(agrupacion="por_tipo", orden_honorarios="paciente")
    wb = load_workbook(BytesIO(build_excel_detalle(a, opciones, enc)))
    filas = list(wb["HONORARIO"].iter_rows())
    textos = [r[0].value for r in filas if r and r[0].value]
    assert "TOTAL PACIENTE DIAZ: 650.00" in textos and "PACIENTE PEREYRA" in textos
    assert not any(str(t).startswith("SUBTOTAL SOCIO") for t in textos)
    subtitulo = next(r[0] for r in filas if r[0].value == "PACIENTE DIAZ")
    assert subtitulo.alignment.horizontal == "centerContinuous" and subtitulo.fill.start_color.rgb.endswith("E0F2FE")
    assert build_pdf_detalle(a, opciones, enc).startswith(b"%PDF")


def test_subtitulos_de_tipo_y_clinica_van_centrados_en_excel():
    a = _armado()
    enc = EncabezadoExport(lineas=["CMC", "OS 411"])
    wb = load_workbook(BytesIO(build_excel_detalle(a, ExportOpciones(agrupacion="por_socio"), enc)))
    celdas = {r[0].value: r[0] for r in wb["Detalle"].iter_rows() if r and r[0].value}
    for texto in ("CONSULTAS", "SANATORIOS", "CLINICA MODELO"):
        assert celdas[texto].alignment.horizontal == "centerContinuous"


def test_ayudante_sin_grupo_se_pega_a_su_cabeza():
    from app.modules.facturacion.equipo import CandidatoEquipo, inferir_equipos

    d = datetime.date(2026, 9, 22)

    def c(id_, med, tipo, *, grupo=None, codigo="126107", clinica=None, paciente="123", fecha=d):
        return CandidatoEquipo(id=id_, cod_medico=med, cod_clinica=clinica, codigo=codigo, tipo_prestador=tipo,
                               paciente=paciente, fecha_practica=fecha, grupo_equipo_id=grupo)

    cands = [
        c(1, "10", "Medico"),
        c(2, "20", "Medico", codigo="999999"),                      # otra cirugía del mismo paciente y día
        c(3, "30", "Ayudante"),                                      # mismo código que la 1 → va con la 1
        c(4, "10", "Ayudante"),                                      # del mismo socio que la 1: no se asiste a sí mismo
        c(5, "40", "Ayudante", paciente="999"),                      # otro paciente: queda suelto
        c(6, "50", "Pediatra", codigo="110401", paciente="777"),
        c(7, "60", "Medico", codigo="110401", paciente="777"),       # cesárea: la cabeza del pediatra
        c(8, "70", "Ayudante", grupo=1),                             # ya tiene grupo: no se toca
    ]
    assert inferir_equipos(cands) == {3: 1, 4: 2, 6: 7}


def test_el_total_de_la_factura_incluye_el_equipo_en_todas_las_agrupaciones():
    """El equipo (ayudante/pediatra) suma SIEMPRE al total de la factura, aunque en Honorarios
    individuales y Sanatorios no sume al socio: RESUMEN GENERAL = suma de todas las filas."""
    filas = _filas_hi()
    esperado = sum((f.subtotal for f in filas), Decimal("0"))
    for agrupacion, extra in (("por_socio", {}), ("por_tipo", {}), ("por_tipo", {"orden_honorarios": "paciente"}),
                              ("plana", {}), ("todo_junto", {})):
        a = armar(filas, ExportOpciones(agrupacion=agrupacion, **extra))
        assert a.resumen.total_general == esperado == Decimal("1665.00"), (agrupacion, extra)
        assert dict(a.resumen.por_tipo)["Honorarios individuales"] == Decimal("1650.00")  # cabezas 1500 + equipo 150
        assert a.total_prestaciones == len(filas)


def test_por_tipo_las_secciones_suman_el_total_general():
    """El subtotal del socio deja afuera al equipo de Honorarios individuales, pero la
    sección lo suma: la suma de las secciones tiene que dar el total general."""
    a = armar(_filas_hi(), ExportOpciones(agrupacion="por_tipo"))
    hi = next(s for s in a.secciones if s.titulo == "HONORARIO")
    assert hi.grupos[0].total_general == Decimal("1500.00")  # socio 10, sin su equipo
    assert hi.total == Decimal("1650.00")  # + ayudante 100 + pediatra 50
    assert sum(s.total for s in a.secciones) == a.resumen.total_general


def test_ayudante_va_con_su_cabeza_aunque_ella_no_se_apunte_a_si_misma():
    """Hay cirugías cargadas sin `grupo_equipo_id` propio y sus ayudantes apuntan a ellas: la
    vista las agrupa, y el export las dejaba separadas (el ayudante como línea de su socio)."""
    cirujano = _fila(1, "10", "ACOSTA, MARIA", "Practica", 5, "RUIZ", monto="1000")        # sin grupo
    ayudante = _fila(2, "30", "BRAVO, LUIS", "Practica", 5, "RUIZ", equipo=1, monto="200")
    ayudante.honorarios, ayudante.ayudante, ayudante.tipo_prestador = Decimal("0"), Decimal("200"), "Ayudante"
    for agrupacion in ("por_socio", "por_tipo", "plana", "todo_junto"):
        a = armar([cirujano, ayudante], ExportOpciones(agrupacion=agrupacion))
        lineas = [l for s in a.secciones for g in s.grupos for l in g.lineas]
        assert [(l.fila.id, [h.id for h in l.hijos]) for l in lineas] == [(1, [2])], agrupacion
        assert a.total_prestaciones == 2

    # Si la cabeza queda afuera del filtro, su equipo también.
    afuera = dataclasses.replace(cirujano, fuera_de_filtro=True)
    a = armar([afuera, ayudante], ExportOpciones(agrupacion="por_socio"))
    assert a.secciones == []

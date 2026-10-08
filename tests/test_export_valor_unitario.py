"""Columna "Valor unitario" del export: honorarios + gastos − coseguro; sin valor para el ayudante."""
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.modules.facturacion.export.armado import armar, valor_unitario
from app.modules.facturacion.export.excel import build_excel_detalle
from app.modules.facturacion.export.pdf import build_pdf_detalle
from app.modules.facturacion.export.schemas import ExportOpciones
from tests.test_export_por_socio import _fila
from app.modules.facturacion.export.encabezado import EncabezadoExport


def test_valor_unitario_resta_el_coseguro_y_el_ayudante_no_lo_lleva():
    f = _fila(1, "10", "ACOSTA, MARIA", "Practica", 5, "RUIZ", monto="100")
    f.honorarios, f.gastos, f.coseguro = Decimal("1000"), Decimal("200"), Decimal("150")
    assert valor_unitario(f) == Decimal("1050.00")
    f.tipo_prestador = "Ayudante"
    assert valor_unitario(f) is None


def test_export_con_la_columna_valor_unitario():
    f = _fila(1, "10", "ACOSTA, MARIA", "Practica", 5, "RUIZ", monto="100")
    f.honorarios, f.gastos, f.coseguro = Decimal("1000"), Decimal("200"), Decimal("150")
    op = ExportOpciones(agrupacion="plana", columnas=["codigo", "honorarios", "coseguro", "valor_unitario"])
    a = armar([f], op)
    enc = EncabezadoExport(lineas=["CMC", "OS 411"])
    wb = load_workbook(BytesIO(build_excel_detalle(a, op, enc)))
    filas = [[c.value for c in r] for r in wb.active.iter_rows()]
    cab = next(r for r in filas if "VALOR UNIT." in r)
    dato = next(r for r in filas if r and r[0] == 1)
    assert dato[cab.index("VALOR UNIT.")] == Decimal("1050.00") or float(dato[cab.index("VALOR UNIT.")]) == 1050.0
    assert build_pdf_detalle(a, op, enc).startswith(b"%PDF")


def test_la_fila_del_ayudante_muestra_lo_que_cobra_en_la_columna_honorarios():
    """El ayudante guarda su monto en `ayudante` con `honorarios` en 0: la columna mostraba 0
    y el SUB. TOTAL sí lo contaba. Ahora muestra lo que cobra, como la vista."""
    cirujano = _fila(1, "10", "ACOSTA, MARIA", "Honorarios individuales", 5, "RUIZ", monto="1000")
    cirujano.tipo_prestador = "Medico"
    ayudante = _fila(2, "30", "BRAVO, LUIS", "Honorarios individuales", 5, "RUIZ", equipo=1, monto="200")
    ayudante.honorarios, ayudante.ayudante, ayudante.tipo_prestador = Decimal("0"), Decimal("200"), "Ayudante"
    cirujano.grupo_equipo_id = 1
    op = ExportOpciones(agrupacion="plana", columnas=["codigo", "honorarios"])
    wb = load_workbook(BytesIO(build_excel_detalle(
        armar([cirujano, ayudante], op), op, EncabezadoExport(lineas=["CMC", "OS 411"]))))
    filas = [[c.value for c in r] for r in wb.active.iter_rows()]
    cab = next(r for r in filas if "HONORARIOS" in r)
    i, sub = cab.index("HONORARIOS"), cab.index("SUB. TOTAL")
    por_id = {r[0]: r for r in filas if r and r[0] in (1, 2)}
    assert float(por_id[1][i]) == 1000.0
    assert float(por_id[2][i]) == 200.0 and float(por_id[2][sub]) == 200.0

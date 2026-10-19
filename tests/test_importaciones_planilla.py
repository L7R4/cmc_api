"""Importación de prestaciones desde la planilla del Colegio.

Servicio sin commit real (`commit` → `flush`, el fixture `db` deshace al final), sobre la
obra social de prueba 888 y el período 209912, que no existen en facturación. Todo en
modo manual (CALCULO M): no hace falta que el código tenga precio.
"""
import datetime
from decimal import Decimal

import pytest
from fastapi import HTTPException
from pydantic import ValidationError
from sqlalchemy import func, select

from app.db.models import Afiliado, DetalleFacturacionCMC, ListadoMedico
from app.modules.importaciones.planilla import servicio
from app.modules.importaciones.planilla.columnas import CLAVES, COLUMNAS, HOJA
from app.modules.importaciones.planilla.plantilla import generar_plantilla
from app.modules.importaciones.planilla.schemas import FilaPlanilla, PlanillaIn

D = Decimal
OS = 888
PERIODO = "209912"
M = DetalleFacturacionCMC


@pytest.fixture
def s(db, monkeypatch):
    monkeypatch.setattr(db, "commit", db.flush)

    async def _no_rollback():
        return None

    monkeypatch.setattr(db, "rollback", _no_rollback)
    return db


async def _medicos(db, n=3) -> list[ListadoMedico]:
    out = (await db.execute(
        select(ListadoMedico).where(
            ListadoMedico.EXISTE == "S", ListadoMedico.es_organizacion == 0, ListadoMedico.MATRICULA_PROV > 0,
        ).order_by(ListadoMedico.NRO_SOCIO).limit(n)
    )).scalars().all()
    assert len(out) == n
    return list(out)


async def _clinica(db) -> ListadoMedico:
    c = (await db.execute(
        select(ListadoMedico).where(ListadoMedico.EXISTE == "S", ListadoMedico.es_organizacion == 1).limit(1)
    )).scalars().first()
    if c is None:
        pytest.skip("No hay una clínica activa.")
    return c


def fila(socio, **kw) -> FilaPlanilla:
    base = dict(
        tipo="CONSULTA", nro_socio=str(socio.NRO_SOCIO), matricula=str(socio.MATRICULA_PROV),
        codigo="420101", cantidad="1", calculo="M", honorarios="24000", fecha="2026-09-08",
        autorizacion="TPL-1", afiliado="PACIENTE DE PRUEBA",
    )
    base.update(kw)
    return FilaPlanilla(**base)


def datos(filas, lote=None) -> PlanillaIn:
    for i, f in enumerate(filas, 2):
        f.fila = i
    return PlanillaIn(obra_social=OS, periodo=PERIODO, archivo="t.xlsx", filas=filas, lote=lote)


async def _cantidad(db) -> int:
    return (await db.execute(
        select(func.count()).select_from(M).where(M.cod_obr == str(OS), M.periodo == PERIODO)
    )).scalar_one()


async def _procesar(db, filas, *, grabar=False, lote=None, cancelado=None):
    return await servicio.procesar(
        db, datos=datos(filas, lote), usuario="test", grabar=grabar, cancelado=cancelado,
    )


def _equipo(a, b, c=None):
    filas = [
        fila(a, tipo="HONORARIOS INDIVIDUALES", codigo="110403", honorarios="420000", equipo="E1", rol="MEDICO"),
        fila(b, tipo="HONORARIOS INDIVIDUALES", codigo="110403", honorarios="", ayudante="105000",
             equipo="E1", rol="AYUDANTE"),
    ]
    if c is not None:
        filas.append(fila(c, tipo="HONORARIOS INDIVIDUALES", codigo="320104", honorarios="121824.98",
                          equipo="E1", rol="PEDIATRA"))
    return filas


# ── Previsualizar ────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_previsualizar_no_escribe_y_arma_el_resumen(s):
    a, b, c = await _medicos(s)
    assert await _cantidad(s) == 0
    out = await _procesar(s, [fila(a), *_equipo(b, c)])

    assert [r.resultado for r in out.filas] == ["grabable"] * 3
    assert out.resumen.grabables == 3 and out.resumen.equipos == 1
    assert out.resumen.importe_total == D("24000.00") + D("420000.00") + D("105000.00")
    assert out.resumen.obra_social == OS and out.resumen.periodo == PERIODO
    assert await _cantidad(s) == 0  # previsualizar no graba nada


@pytest.mark.asyncio
async def test_confirmar_graba_el_equipo_enlazado_y_marca_el_lote(s):
    a, b, c = await _medicos(s)
    out = await _procesar(s, [fila(a), *_equipo(b, c, a)], grabar=True, lote="abc12345-0001")

    assert out.resumen.grabables == 4 and all(r.resultado == "grabada" for r in out.filas)
    assert await _cantidad(s) == 4

    ids = [r.id_detalle for r in out.filas]
    individual, cirujano, ayudante, pediatra = [await s.get(M, i) for i in ids]
    assert individual.grupo_equipo_id is None
    # El cirujano es la cabeza: todos los integrantes apuntan a él y llevan su tipo.
    assert cirujano.grupo_equipo_id == cirujano.id_detalle_prestaciones
    assert ayudante.grupo_equipo_id == pediatra.grupo_equipo_id == cirujano.id_detalle_prestaciones
    assert {cirujano.tipo, ayudante.tipo, pediatra.tipo} == {"Honorarios individuales"}
    assert ayudante.ayudante == D("105000.00") and ayudante.honorarios == 0
    assert pediatra.tpo_funcion == "P"
    # La previsualización mostró el tipo con el que quedaron.
    assert [r.tipo for r in out.filas[1:]] == ["Honorarios individuales"] * 3
    assert all(r.validacion_detalle.startswith("Importado de planilla · lote abc12345-0001")
               for r in (individual, cirujano, ayudante, pediatra))

    lote = await servicio.estado_lote(s, obra_social=OS, periodo=PERIODO, lote="abc12345-0001")
    assert lote.grabadas == 4
    assert (await servicio.estado_lote(s, obra_social=OS, periodo=PERIODO, lote="abc12345-0002")).grabadas == 0


@pytest.mark.asyncio
async def test_si_falla_una_fila_del_equipo_no_entra_ninguna(s):
    a, b, _ = await _medicos(s)
    malas = _equipo(a, b)
    malas[1].nro_socio = "999999999"  # el ayudante no existe
    malas[1].matricula = ""
    out = await _procesar(s, [fila(a, autorizacion="TPL-9"), *malas])

    assert out.filas[0].resultado == "grabable"
    assert [r.resultado for r in out.filas[1:]] == ["omitida", "omitida"]
    assert "No existe el socio 999999999" in out.filas[2].motivo
    assert "no entra" in out.filas[1].motivo
    assert out.resumen.grabables == 1 and out.resumen.omitidas == 2


@pytest.mark.asyncio
async def test_matricula_distinta_pide_elegir_y_bloquea_confirmar(s):
    a, b, _ = await _medicos(s)
    f = fila(a, matricula=str(b.MATRICULA_PROV))
    out = await _procesar(s, [f])
    r = out.filas[0]
    assert r.resultado == "elegir_socio" and out.resumen.por_elegir == 1
    assert a.NRO_SOCIO in [c.nro_socio for c in r.candidatos]

    with pytest.raises(HTTPException) as e:
        await _procesar(s, [f], grabar=True)
    assert e.value.status_code == 422 and "sin socio" in e.value.detail

    # Con el socio elegido la fila entra (el operador decidió).
    f.nro_socio_elegido = a.NRO_SOCIO
    out = await _procesar(s, [f])
    assert out.filas[0].resultado == "grabable" and out.filas[0].nro_socio == a.NRO_SOCIO


@pytest.mark.asyncio
async def test_descartada_no_se_procesa_ni_bloquea(s):
    a, b, _ = await _medicos(s)
    mala = fila(b, tipo="XX")
    mala.descartada = True
    out = await _procesar(s, [fila(a), mala], grabar=True)
    assert [r.resultado for r in out.filas] == ["grabada", "descartada"]
    assert out.resumen.descartadas == 1 and await _cantidad(s) == 1


@pytest.mark.asyncio
async def test_reimportar_marca_las_filas_como_ya_cargadas(s):
    a, b, c = await _medicos(s)
    filas = [fila(a), *_equipo(b, c)]
    await _procesar(s, filas, grabar=True)

    out = await _procesar(s, [fila(a), *_equipo(b, c)])
    assert [r.resultado for r in out.filas] == ["duplicada"] * 3
    assert "Ya está cargada" in out.filas[0].motivo
    assert out.resumen.duplicadas == 3 and out.resumen.grabables == 0

    # Sin autorización ni fecha no hay cómo reconocerla: entra y avisa.
    sin = fila(a, autorizacion="", fecha="")
    out = await _procesar(s, [sin], grabar=True)
    out = await _procesar(s, [fila(a, autorizacion="", fecha="")])
    assert out.filas[0].resultado == "grabable" and "no se puede detectar" in out.filas[0].aviso


@pytest.mark.asyncio
async def test_sanatorio_lleva_la_clinica_y_el_medico_ejecutor(s):
    a, _, _ = await _medicos(s)
    clin = await _clinica(s)
    out = await _procesar(s, [fila(a, tipo="SANATORIO", clinica=str(clin.NRO_SOCIO))], grabar=True)
    row = await s.get(M, out.filas[0].id_detalle)
    assert row.tipo == "Sanatorio" and row.cod_clinica == clin.NRO_SOCIO and str(row.cod_med) == str(a.NRO_SOCIO)

    out = await _procesar(s, [fila(a, tipo="SANATORIO", clinica=str(a.NRO_SOCIO), autorizacion="X")])
    assert "no es una clínica" in out.filas[0].motivo


async def _afiliado(db, dni):
    return (await db.execute(select(Afiliado).where(Afiliado.dni == dni))).scalars().first()


@pytest.mark.asyncio
async def test_el_paciente_que_no_esta_en_el_padron_se_registra_al_importar(s):
    a, _, _ = await _medicos(s)
    dni = "TPL-AF-0001"
    assert await _afiliado(s, dni) is None
    f = fila(a, nro_afiliado=dni, afiliado="PEREZ, JUAN DE PRUEBA")

    # Previsualizar avisa pero no crea nada.
    out = await _procesar(s, [f])
    assert out.filas[0].resultado == "grabable" and "se registra al importar" in out.filas[0].aviso
    assert out.resumen.afiliados_nuevos == 1
    assert await _afiliado(s, dni) is None

    # Confirmar lo registra con el nombre de la planilla y lo copia a la prestación.
    out = await _procesar(s, [f], grabar=True)
    af = await _afiliado(s, dni)
    assert af is not None and af.nombre == "PEREZ, JUAN DE PRUEBA" and af.usuario == "test"
    row = await s.get(M, out.filas[0].id_detalle)
    assert row.dni_p == dni and row.nom_ape_p == "PEREZ, JUAN DE PRUEBA"

    # Ya registrado, no se vuelve a crear ni a avisar.
    out = await _procesar(s, [fila(a, nro_afiliado=dni, afiliado="OTRO NOMBRE", autorizacion="TPL-2")])
    assert out.resumen.afiliados_nuevos == 0 and "padrón" not in out.filas[0].aviso


@pytest.mark.asyncio
async def test_varias_filas_del_mismo_paciente_nuevo_lo_registran_una_sola_vez(s):
    a, b, _ = await _medicos(s)
    dni = "TPL-AF-0002"
    filas = [fila(a, nro_afiliado=dni, afiliado="", autorizacion="A1"), fila(b, nro_afiliado=dni, afiliado="GOMEZ, ANA", autorizacion="A2")]
    out = await _procesar(s, filas, grabar=True)
    assert out.resumen.afiliados_nuevos == 1
    assert (await _afiliado(s, dni)).nombre == "GOMEZ, ANA"  # el nombre de la primera fila que lo trae
    n = (await s.execute(select(func.count()).select_from(Afiliado).where(Afiliado.dni == dni))).scalar_one()
    assert n == 1


@pytest.mark.asyncio
async def test_si_la_fila_no_entra_no_se_registra_su_paciente(s):
    a, _, _ = await _medicos(s)
    dni = "TPL-AF-0003"
    out = await _procesar(s, [fila(a, nro_afiliado=dni, afiliado="NO ENTRA", cantidad="0")], grabar=False)
    assert out.filas[0].resultado == "omitida" and out.resumen.afiliados_nuevos == 0
    with pytest.raises(HTTPException):  # nada para importar
        await _procesar(s, [fila(a, nro_afiliado=dni, afiliado="NO ENTRA", cantidad="0")], grabar=True)
    assert await _afiliado(s, dni) is None


@pytest.mark.asyncio
async def test_el_resumen_cuenta_lo_que_no_se_puede_verificar_como_repetido(s):
    a, b, c = await _medicos(s)
    out = await _procesar(s, [
        fila(a),                                              # con autorización y fecha
        fila(b, autorizacion=""),                             # sin autorización, con fecha
        fila(c, autorizacion="", fecha=""),                   # sin autorización ni fecha
        fila(a, autorizacion="", fecha="", cantidad="0"),     # no entra: no cuenta
    ])
    assert out.resumen.grabables == 3
    assert out.resumen.sin_autorizacion == 2 and out.resumen.no_verificables == 1


@pytest.mark.parametrize("cambios, esperado", [
    ({"tipo": "OTRO"}, "TIPO"),
    ({"tipo": ""}, "Falta el TIPO"),
    ({"cantidad": "0"}, "CANTIDAD"),
    ({"cantidad": ""}, "Falta la CANTIDAD"),
    ({"cantidad": "1,5"}, "entero"),
    ({"fecha": "32/08/2026"}, "FECHA"),
    ({"fecha": "ayer"}, "FECHA"),
    ({"codigo": ""}, "Falta el CODIGO"),
    ({"nro_socio": ""}, "Falta el NRO_SOCIO"),
    ({"nro_socio": "abc"}, "NRO_SOCIO"),
    ({"via": "X"}, "VIA"),
    ({"calculo": "Z"}, "CALCULO"),
    ({"honorarios": "-5"}, "negativo"),
    ({"honorarios": "mucho"}, "importe"),
    ({"honorarios": "", "gastos": "", "ayudante": ""}, "ningún importe"),
    ({"clinica": "9882"}, "CLINICA"),                       # en CONSULTA no va clínica
    ({"rol": "AYUDANTE"}, "necesita un EQUIPO"),
    ({"rol": "JEFE"}, "ROL"),
    ({"honorarios": "10", "ayudante": "5"}, "AYUDANTE junto con"),
    ({"autorizacion": "x" * 31}, "AUTORIZACION"),
    ({"nro_afiliado": "1" * 21}, "NRO_AFILIADO"),
    ({"codigo": "5000"}, "Código no permitido"),
    ({"nro_socio": "999999999", "matricula": ""}, "No existe el socio"),
])
@pytest.mark.asyncio
async def test_cada_error_de_la_planilla_se_informa_con_su_motivo(s, cambios, esperado):
    a, _, _ = await _medicos(s)
    out = await _procesar(s, [fila(a, **cambios)])
    r = out.filas[0]
    assert r.resultado == "omitida", r.motivo
    assert esperado in r.motivo


@pytest.mark.asyncio
async def test_equipo_sin_cirujano_o_con_codigos_distintos_no_entra(s):
    a, b, _ = await _medicos(s)
    solo_ayudantes = [f for f in _equipo(a, b) if f.rol == "AYUDANTE"] * 2
    out = await _procesar(s, solo_ayudantes)
    assert all(r.resultado == "omitida" for r in out.filas) and "exactamente una fila con ROL MEDICO" in out.filas[0].motivo

    otro = _equipo(a, b)
    otro[1].codigo = "110404"
    out = await _procesar(s, otro)
    assert all(r.resultado == "omitida" for r in out.filas) and "mismo CODIGO" in out.filas[0].motivo


@pytest.mark.asyncio
async def test_periodos_informa_el_minimo_para_cargar(s):
    # Sin ningún período cerrado no hay piso.
    sin = await servicio.periodos(s, OS)
    assert sin.minimo is None

    # Con períodos cerrados el piso es el último cerrado + 1, y es el sugerido.
    from app.db.models import FacturacionCMC
    cerrada = (await s.execute(
        select(FacturacionCMC.cod_obr).where(FacturacionCMC.estado.in_(("C", "L", "LC"))).limit(1)
    )).scalar_one_or_none()
    if cerrada is None:
        pytest.skip("No hay ninguna factura cerrada.")
    con = await servicio.periodos(s, int(cerrada))
    assert con.minimo == con.sugerido and con.minimo is not None
    assert all(p.periodo < con.minimo for p in con.periodos if p.cerrado)


@pytest.mark.asyncio
async def test_obra_social_y_periodo_se_validan(s):
    a, _, _ = await _medicos(s)
    with pytest.raises(HTTPException) as e:
        await servicio.procesar(
            s, datos=PlanillaIn(obra_social=99999, periodo=PERIODO, filas=[fila(a)]), usuario="t", grabar=False)
    assert e.value.status_code == 422 and "no existe" in e.value.detail

    with pytest.raises(ValidationError):  # el período es AAAAMM
        PlanillaIn(obra_social=OS, periodo="20261", filas=[fila(a)])


# ── Todo o nada ──────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_cancelar_antes_del_commit_no_hace_commit(db, monkeypatch):
    a, b, c = await _medicos(db)
    commits = []

    async def espia():
        commits.append(1)

    monkeypatch.setattr(db, "commit", espia)

    async def cancelado():
        return True

    with pytest.raises(servicio.ImportacionCancelada):
        await _procesar(db, [fila(a), *_equipo(b, c)], grabar=True, cancelado=cancelado)
    assert commits == []


@pytest.mark.asyncio
async def test_cancelar_justo_antes_del_commit_tambien_lo_evita(db, monkeypatch):
    """El operador cancela cuando ya se agregaron todas las filas: la última revisión,
    inmediatamente antes del commit, tiene que frenarlo igual."""
    a, _, _ = await _medicos(db)
    commits = []

    async def espia():
        commits.append(1)

    monkeypatch.setattr(db, "commit", espia)
    llamadas = []

    async def cancelado():
        llamadas.append(1)
        return len(llamadas) > 1  # la 1.ª (al empezar el grabado) no; la última (pre-commit) sí

    with pytest.raises(servicio.ImportacionCancelada):
        await _procesar(db, [fila(a)], grabar=True, cancelado=cancelado)
    assert commits == [] and len(llamadas) == 2


@pytest.mark.asyncio
async def test_un_error_al_grabar_no_deja_commit(db, monkeypatch):
    a, b, _ = await _medicos(db)
    commits = []

    async def espia():
        commits.append(1)

    monkeypatch.setattr(db, "commit", espia)
    from app.modules.facturacion import service as fact

    async def falla(*a, **k):
        raise RuntimeError("se cayó la base")

    monkeypatch.setattr(fact, "_ensure_factura_abierta", falla)
    with pytest.raises(RuntimeError):
        await _procesar(db, [fila(a), fila(b)], grabar=True)
    assert commits == []


# ── Plantilla ────────────────────────────────────────────────────────────────


def test_la_plantilla_sale_de_las_columnas_del_servicio(tmp_path):
    from openpyxl import load_workbook

    p = tmp_path / "plantilla.xlsx"
    p.write_bytes(generar_plantilla())
    wb = load_workbook(p)
    assert wb.sheetnames == [HOJA, "Instrucciones"]
    encabezados = [c.value for c in wb[HOJA][1]]
    assert encabezados == [c.clave + (" *" if c.obligatoria else "") for c in COLUMNAS]
    # Las columnas que se sacaron no vuelven por error.
    assert not {"OBRA_SOCIAL", "PERIODO", "PRESTADOR", "REFERENCIA"} & set(CLAVES)
    assert "COMENTARIOS" in CLAVES
    # La planilla (hoja de datos) está vacía.
    assert wb[HOJA].max_row == 500 and all(c.value is None for c in wb[HOJA][2])
